"""Generate controlled fixtures; no real course material is bundled."""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import av
import numpy as np
from tests.test_documents import make_epub, make_mixed_pdf
from docx import Document


def main():
    root = Path(__file__).resolve().parent.parent / 'tmp' / 'fixtures'
    root.mkdir(parents=True, exist_ok=True)
    docs = root / 'documents'; docs.mkdir(exist_ok=True)
    make_epub(docs / 'chapters.epub')
    make_mixed_pdf(docs / 'mixed.pdf')
    doc = Document(); doc.add_heading('Course Introduction', level=1)
    doc.add_paragraph('Keep the original files unchanged.'); doc.save(docs / 'notes.docx')
    (docs / 'notes.txt').write_text('Original source text. Café.\nSecond paragraph.', encoding='utf-8')
    (docs / 'notes.md').write_text('# A Markdown heading\n\nPreserve **source** wording.', encoding='utf-8')
    audio = root / 'media' / 'lesson.wav'
    if not audio.exists():
        raise SystemExit('Generate lesson.wav with scripts/make_speech.ps1 first.')
    for extension, codec in [('mp4', 'aac'), ('mkv', 'aac'), ('mov', 'aac'), ('m4a', 'aac'), ('mp3', 'libmp3lame'), ('flac', 'flac'), ('ts', 'aac')]:
        target = audio.with_suffix('.' + extension)
        with av.open(str(audio)) as source, av.open(str(target), 'w') as output:
            rate = 44100
            stream = output.add_stream(codec, rate=rate)
            stream.layout = 'mono'
            video = None
            video_rate = 25 if extension == 'ts' else 1
            if extension in {'mp4', 'mkv', 'mov', 'ts'}:
                video = output.add_stream('libx264' if extension == 'ts' else 'mpeg4', rate=video_rate)
                video.width, video.height, video.pix_fmt = 320, 180, 'yuv420p'
            sample_format = 's16' if codec == 'flac' else 'fltp'
            resampler = av.AudioResampler(format=sample_format, layout='mono', rate=rate)
            count = 0
            for frame in source.decode(audio=0):
                for resampled in resampler.resample(frame):
                    count += resampled.samples
                    for packet in stream.encode(resampled):
                        output.mux(packet)
            for resampled in resampler.resample(None):
                count += resampled.samples
                for packet in stream.encode(resampled):
                    output.mux(packet)
            for packet in stream.encode(None):
                output.mux(packet)
            if video is not None:
                for i in range(int(count / rate * video_rate) + 1):
                    pixels = np.full((180, 320, 3), 30 + i % 100, dtype=np.uint8)
                    frame = av.VideoFrame.from_ndarray(pixels, format='rgb24')
                    frame.pts = i
                    for packet in video.encode(frame):
                        output.mux(packet)
                for packet in video.encode(None):
                    output.mux(packet)
        print(target)
    with __import__('pymupdf').open(docs / 'mixed.pdf') as pdf:
        pdf[2].get_pixmap().save(str(root / 'mixed-preview.png'))


if __name__ == '__main__':
    main()
