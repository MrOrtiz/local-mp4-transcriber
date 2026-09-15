from dataclasses import replace
from pathlib import Path
import threading
import zipfile
import pytest
import pymupdf
from docx import Document
from ebooklib import epub

from course_transcriber.documents import extract_document, tessdata_path
from course_transcriber.types import Settings, Cancelled


def extract(path, settings=Settings()):
    return extract_document(path, settings, threading.Event(), lambda _: None)


@pytest.mark.parametrize('suffix,encoding', [('.txt', 'utf-8'), ('.md', 'utf-8-sig'), ('.markdown', 'utf-16')])
def test_text_unicode_preserved(tmp_path, suffix, encoding):
    path = tmp_path / ('lesson' + suffix)
    original = '# Chapter\n\nCafé — naïve.\n1. Exact source words.'
    path.write_bytes(original.encode(encoding))
    assert extract(path).text == original


def test_docx_paragraph_table_order(tmp_path):
    path = tmp_path / 'lesson.docx'
    doc = Document(); doc.add_heading('Heading One', level=1)
    doc.add_paragraph('Before table')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Cell alpha'; table.cell(0, 1).text = 'Cell beta'
    doc.add_paragraph('After table'); doc.save(path)
    result = extract(path)
    assert '# Heading One' in result.text
    assert result.text.index('Before table') < result.text.index('Cell alpha') < result.text.index('After table')


def test_pdf_selectable_page_references(tmp_path):
    path = tmp_path / 'lesson.pdf'
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), 'First page heading')
    doc.new_page().insert_text((72, 72), 'Second page content')
    doc.save(path); doc.close()
    result = extract(path)
    assert result.metadata['pages'] == 2
    assert '[Page 1]\nFirst page heading' in result.text
    assert '[Page 2]\nSecond page content' in result.text


def make_mixed_pdf(path):
    scan = pymupdf.open()
    page = scan.new_page(width=500, height=180)
    page.insert_text((25, 60), 'Scanned lesson content', fontsize=24)
    page.insert_text((25, 105), 'Preserve these source words.', fontsize=20)
    png = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes('png')
    doc = pymupdf.open()
    doc.new_page().insert_text((40, 60), 'Selectable first page', fontsize=18)
    second = doc.new_page()
    second.insert_image(pymupdf.Rect(40, 100, 540, 280), stream=png)
    third = doc.new_page()
    third.insert_text((40, 60), 'Selectable mixed heading', fontsize=18)
    third.insert_image(pymupdf.Rect(40, 100, 540, 280), stream=png)
    doc.save(path); doc.close(); scan.close()


def test_pdf_scanned_and_mixed_ocr(tmp_path):
    if tessdata_path() is None:
        pytest.skip('English OCR data not installed; run install-ocr.cmd')
    path = tmp_path / 'mixed.pdf'
    make_mixed_pdf(path)
    result = extract(path)
    assert result.metadata['pages'] == 3
    assert 'Selectable first page' in result.text
    assert 'Selectable mixed heading' in result.text
    assert result.text.count('Scanned lesson content') == 2
    assert result.text.index('Selectable mixed heading') < result.text.rindex('Scanned lesson content')
    assert sum('local OCR used' in note for note in result.warnings) == 2


def test_missing_ocr_fails_instead_of_claiming_complete(tmp_path, monkeypatch):
    from course_transcriber import documents
    path = tmp_path / 'mixed.pdf'; make_mixed_pdf(path)
    monkeypatch.setattr(documents, 'tessdata_path', lambda _: None)
    with pytest.raises(RuntimeError, match='requires OCR'):
        extract(path)


def test_protected_pdf(tmp_path):
    path = tmp_path / 'protected.pdf'
    doc = pymupdf.open(); doc.new_page().insert_text((72, 72), 'Secret')
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='password')
    doc.close()
    with pytest.raises(ValueError, match='Protected PDF'):
        extract(path)


def make_epub(path):
    book = epub.EpubBook(); book.set_identifier('fixture'); book.set_title('Fixture Book'); book.set_language('en')
    one = epub.EpubHtml(title='One', file_name='one.xhtml', lang='en')
    one.content = '<h1>Chapter One</h1><p>First <b>important</b> sentence.</p>'
    two = epub.EpubHtml(title='Two', file_name='two.xhtml', lang='en')
    two.content = '<h1>Chapter Two</h1><p>Second chapter content.</p>'
    book.add_item(two); book.add_item(one)  # Deliberately reverse archive order.
    book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav())
    book.spine = [one, two]; book.toc = [one, two]
    epub.write_epub(str(path), book)


def test_epub_spine_reading_order(tmp_path):
    path = tmp_path / 'lesson.epub'; make_epub(path)
    result = extract(path)
    assert result.text.index('Chapter One') < result.text.index('Chapter Two')
    assert 'First important sentence.' in result.text
    assert result.metadata['language'] == 'en'


def test_protected_epub(tmp_path):
    path = tmp_path / 'protected.epub'; make_epub(path)
    with zipfile.ZipFile(path, 'a') as archive:
        archive.writestr('META-INF/encryption.xml', '<encryption><EncryptionMethod Algorithm="aes"/></encryption>')
    with pytest.raises(ValueError, match='Protected EPUB'):
        extract(path)


def test_mobi_missing_dependency(tmp_path, monkeypatch):
    from course_transcriber import documents
    path = tmp_path / 'lesson.mobi'; path.write_bytes(b'dummy')
    monkeypatch.setattr(documents, 'converter_path', lambda _: None)
    with pytest.raises(RuntimeError, match='requires Calibre'):
        extract(path)


def test_cancel_before_document(tmp_path):
    path = tmp_path / 'lesson.txt'; path.write_text('Hello')
    event = threading.Event(); event.set()
    with pytest.raises(Cancelled):
        extract_document(path, Settings(), event, lambda _: None)


def test_real_mobi_conversion(tmp_path):
    import subprocess
    from course_transcriber.documents import converter_path
    exe = converter_path()
    if not exe:
        pytest.skip('Calibre is not installed')
    source = tmp_path / 'fixture.epub'; make_epub(source)
    mobi = tmp_path / 'fixture.mobi'
    subprocess.run([exe, str(source), str(mobi)], check=True, capture_output=True, timeout=120,
                   creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    result = extract(mobi)
    assert 'Chapter One' in result.text and 'Second chapter content.' in result.text
    assert result.text.index('Chapter One') < result.text.index('Chapter Two')


def test_corrupt_mobi_reports_failure(tmp_path):
    from course_transcriber.documents import converter_path
    if not converter_path():
        pytest.skip('Calibre is not installed')
    path = tmp_path / 'corrupt.mobi'; path.write_bytes(b'Not a valid ebook')
    with pytest.raises(ValueError, match='MOBI conversion failed'):
        extract(path)
